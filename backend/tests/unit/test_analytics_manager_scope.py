"""HRP-641: analytics and the XLSX export answer for the caller's subtree.

Seven routes called the service with a tenant id and nothing else, so the
head of a three-person department read the whole company's aggregates and
exported every employee's assessments to a spreadsheet. The lists those
figures summarise had been scoped since HRP-626, which made the boundary
arithmetic: company total minus own total.
"""

from __future__ import annotations

import inspect
import io
import uuid
from datetime import date, datetime, timezone

import pytest_asyncio
from app.core.security import create_access_token, hash_password
from app.modules.analytics.service import _ai_summary_cache_key
from app.modules.analytics.tasks import export_assessments_task
from app.modules.assessment.models import CPA, PDP, Assessment, AssessmentResult
from app.modules.auth.models import Role, User, user_roles
from app.modules.company.models import Division
from app.modules.competence.models import Competence, CompetenceGroup
from app.modules.employee.models import Employee
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload


async def _user(db: AsyncSession, tenant, code: str) -> User:
    result = await db.execute(select(Role).where(Role.code == code))
    role = result.scalars().first()
    if role is None:
        role = Role(name=code.title(), code=code, is_system=True)
        db.add(role)
        await db.commit()
        await db.refresh(role)
    u = User(
        email=f"{code}-{uuid.uuid4().hex[:8]}@test.com",
        password_hash=hash_password("x"),
        first_name=code,
        last_name="X",
        tenant_id=tenant.id,
        email_verified_at=datetime.now(timezone.utc),
    )
    db.add(u)
    await db.commit()
    await db.refresh(u)
    await db.execute(user_roles.insert().values(user_id=u.id, role_id=role.id))
    await db.commit()
    db.expunge(u)
    result = await db.execute(
        select(User).options(selectinload(User.roles)).where(User.id == u.id)
    )
    return result.scalar_one()


def _headers(u: User) -> dict[str, str]:
    return {
        "Authorization": f"Bearer {create_access_token(str(u.id), str(u.tenant_id))}"
    }


def _rows(content: bytes) -> list[tuple]:
    sheet = load_workbook(io.BytesIO(content)).active
    return list(sheet.iter_rows(min_row=2, values_only=True))


@pytest_asyncio.fixture
async def reports(
    db: AsyncSession, tenant, assessment_statuses, assessment_types
) -> dict:
    """One assessee, one plan and one CPA score per division: 1 vs 2."""
    mgr_user = await _user(db, tenant, "manager")
    lonely_user = await _user(db, tenant, "manager")
    admin_user = await _user(db, tenant, "admin")

    mgr_emp = Employee(
        user_id=mgr_user.id, tenant_id=tenant.id, hire_date=date(2024, 1, 1)
    )
    lonely_emp = Employee(
        user_id=lonely_user.id, tenant_id=tenant.id, hire_date=date(2024, 1, 1)
    )
    db.add_all([mgr_emp, lonely_emp])
    await db.commit()
    await db.refresh(mgr_emp)
    await db.refresh(lonely_emp)

    mine = Division(
        tenant_id=tenant.id, name=f"Mine {uuid.uuid4().hex[:4]}", manager_id=mgr_emp.id
    )
    theirs = Division(tenant_id=tenant.id, name=f"Theirs {uuid.uuid4().hex[:4]}")
    db.add_all([mine, theirs])
    await db.commit()
    await db.refresh(mine)
    await db.refresh(theirs)

    group = CompetenceGroup(title=f"G-{uuid.uuid4().hex[:6]}", tenant_id=tenant.id)
    db.add(group)
    await db.commit()
    await db.refresh(group)
    comp = Competence(
        title=f"C-{uuid.uuid4().hex[:6]}", group_id=group.id, tenant_id=tenant.id
    )
    db.add(comp)
    await db.commit()
    await db.refresh(comp)

    cpa = CPA(
        tenant_id=tenant.id,
        title=f"Round {uuid.uuid4().hex[:4]}",
        author_id=admin_user.id,
        type_id=assessment_types["360"].id,
    )
    db.add(cpa)
    await db.commit()
    await db.refresh(cpa)
    cpa_id = cpa.id

    out: dict = {
        "mgr_user": mgr_user,
        "mgr_employee": mgr_emp,
        "lonely_user": lonely_user,
        "admin_user": admin_user,
        "cpa_id": cpa_id,
    }
    for key, division in (("mine", mine), ("theirs", theirs)):
        emp_user = await _user(db, tenant, "employee")
        emp = Employee(
            user_id=emp_user.id,
            tenant_id=tenant.id,
            division_id=division.id,
            hire_date=date(2024, 1, 1),
        )
        db.add(emp)
        await db.commit()
        await db.refresh(emp)

        assessment = Assessment(
            tenant_id=tenant.id,
            title=f"{key} assessment",
            employee_id=emp.id,
            initiator_id=admin_user.id,
            type_id=assessment_types["360"].id,
            status_id=assessment_statuses["draft"].id,
            cpa_id=cpa_id,
        )
        pdp = PDP(
            tenant_id=tenant.id,
            title=f"{key} plan",
            employee_id=emp.id,
            author_id=admin_user.id,
            status="draft",
        )
        db.add_all([assessment, pdp])
        await db.commit()
        await db.refresh(assessment)
        await db.refresh(pdp)

        db.add(
            AssessmentResult(
                assessment_id=assessment.id, competence_id=comp.id, avg_score=4.0
            )
        )
        await db.commit()

        out[f"{key}_employee"] = emp
        out[f"{key}_assessment"] = assessment
        out[f"{key}_pdp"] = pdp
    return out


class TestAggregatesFollowTheSubtree:
    async def test_manager_counts_only_their_own_people(
        self, client: AsyncClient, reports
    ):
        h = _headers(reports["mgr_user"])
        assert (await client.get("/api/analytics/assessments", headers=h)).json()[
            "total"
        ] == 1
        assert (await client.get("/api/analytics/pdp", headers=h)).json()["total"] == 1

    async def test_admin_still_counts_the_workspace(self, client: AsyncClient, reports):
        h = _headers(reports["admin_user"])
        assert (await client.get("/api/analytics/assessments", headers=h)).json()[
            "total"
        ] == 2
        assert (await client.get("/api/analytics/pdp", headers=h)).json()["total"] == 2

    async def test_a_manager_of_nobody_gets_zero_not_everything(
        self, client: AsyncClient, reports
    ):
        """The row where ``if not visible_ids`` degrades into "no filter"."""
        h = _headers(reports["lonely_user"])
        assert (await client.get("/api/analytics/assessments", headers=h)).json()[
            "total"
        ] == 0
        assert (await client.get("/api/analytics/pdp", headers=h)).json()["total"] == 0

    async def test_cpa_comparison_is_scoped_by_assessee(
        self, client: AsyncClient, reports
    ):
        params = {
            "cpa_id_1": str(reports["cpa_id"]),
            "cpa_id_2": str(reports["cpa_id"]),
        }
        manager = await client.get(
            "/api/analytics/cpa-comparison",
            headers=_headers(reports["mgr_user"]),
            params=params,
        )
        assert manager.status_code == 200, manager.text
        assert [c["employee_id"] for c in manager.json()["comparisons"]] == [
            str(reports["mine_employee"].id)
        ]

        admin = await client.get(
            "/api/analytics/cpa-comparison",
            headers=_headers(reports["admin_user"]),
            params=params,
        )
        assert len(admin.json()["comparisons"]) == 2


class TestExportFollowsTheSubtree:
    async def test_the_spreadsheet_holds_only_the_managed_subtree(
        self, client: AsyncClient, reports
    ):
        resp = await client.post(
            "/api/analytics/export/assessments", headers=_headers(reports["mgr_user"])
        )
        assert resp.status_code == 200
        assert [r[2] for r in _rows(resp.content)] == [str(reports["mine_employee"].id)]

    async def test_a_manager_of_nobody_exports_an_empty_sheet(
        self, client: AsyncClient, reports
    ):
        resp = await client.post(
            "/api/analytics/export/assessments",
            headers=_headers(reports["lonely_user"]),
        )
        assert resp.status_code == 200
        assert _rows(resp.content) == []

    async def test_the_background_export_is_given_the_scope_not_the_tenant(
        self, client: AsyncClient, reports, monkeypatch
    ):
        """Resolving the scope inside the worker would resolve it for the
        tenant, not for whoever asked."""
        captured: list[tuple] = []

        def _fake_enqueue(task, *args, **kwargs):
            captured.append(args)

            class _Result:
                id = "task-1"

            return _Result()

        monkeypatch.setattr(
            "app.core.task_enqueue.enqueue_task", _fake_enqueue, raising=True
        )
        resp = await client.post(
            "/api/analytics/export/assessments/async",
            headers=_headers(reports["mgr_user"]),
        )
        assert resp.status_code == 200, resp.text
        assert len(captured) == 1
        _tenant_id, employee_ids = captured[0]
        assert set(employee_ids) == {
            str(reports["mine_employee"].id),
            str(reports["mgr_employee"].id),
        }

    def test_the_worker_signature_refuses_to_default_to_unscoped(self):
        """A message queued by the previous release fails loudly rather
        than quietly exporting the workspace."""
        params = inspect.signature(export_assessments_task).parameters
        assert params["employee_ids"].default is inspect.Parameter.empty


def test_the_ai_summary_cache_key_separates_scopes():
    """A guessed fingerprint must not hand a manager the admin's summary."""
    tenant_id = uuid.uuid4()
    fingerprint = "same-data-version"
    admin, manager = uuid.uuid4(), uuid.uuid4()
    assert _ai_summary_cache_key(
        tenant_id, "en", fingerprint, user_id=admin
    ) != _ai_summary_cache_key(tenant_id, "en", fingerprint, user_id=manager)
