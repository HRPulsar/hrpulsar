"""HRP-744: the canvas export must be what the canvas is showing.

The workbook ignored the toolbar entirely — every candidate, both Manager
and AI rows, the default round and the default scale, whatever the
recruiter had filtered down to on screen. The CSV honoured the filters
(it was assembled in the browser from the already-filtered rows) but had
no Total column, because that second copy of the rendering rules never
grew one. Both formats read one builder now, so they cannot disagree
again.
"""

from __future__ import annotations

import csv
import io
import re
import uuid
from pathlib import Path

import pytest
from app.main import app
from app.modules.recruitment import service
from app.modules.recruitment.report_xlsx import (
    build_canvas_rows,
    render_canvas_csv,
    render_canvas_xlsx,
)
from app.modules.recruitment.schemas import (
    CandidateCreate,
    CandidateVacancyCreate,
    VacancyCreate,
    VacancyProfileUpdate,
)
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

_COMPETENCES = [
    {"id": "python-skills", "name": "Python", "group": "Hard"},
    {"id": "communication", "name": "Communication", "group": "Soft"},
]


async def _vacancy_with_two_candidates(db: AsyncSession, tenant, user):
    vacancy = await service.create_vacancy(
        db, tenant.id, user.id, VacancyCreate(title=f"V {uuid.uuid4().hex[:5]}")
    )
    vacancy_id = uuid.UUID(str(vacancy["id"]))
    await service.save_profile(
        db,
        tenant.id,
        vacancy_id,
        VacancyProfileUpdate(profile_data={"competences": _COMPETENCES}),
    )
    cv_ids: list[uuid.UUID] = []
    for name in ("Ekaterina", "Maria"):
        cand = await service.create_candidate(
            db,
            tenant.id,
            user.id,
            CandidateCreate(
                first_name=name,
                last_name="Ivanova",
                email=f"{uuid.uuid4().hex[:8]}@example.com",
            ),
        )
        cv = await service.attach_candidate(
            db,
            tenant.id,
            user.id,
            CandidateVacancyCreate(
                candidate_id=uuid.UUID(str(cand["id"])), vacancy_id=vacancy_id
            ),
        )
        cv_ids.append(uuid.UUID(str(cv["id"])))
    return vacancy_id, cv_ids


def _matrix_stub() -> dict:
    """A matrix payload with both halves filled, one divergent cell."""
    first, second = str(uuid.uuid4()), str(uuid.uuid4())
    comp_a, comp_b = str(uuid.uuid4()), str(uuid.uuid4())
    return {
        "max_score": 4.0,
        "scale_name": "Standard 1-4",
        "round": "latest",
        "divergence_threshold": 1.0,
        "competences": [
            {"id": comp_a, "name": "Python"},
            {"id": comp_b, "name": "Communication"},
        ],
        "candidates": [
            {
                "candidate_vacancy_id": first,
                "name": "Ekaterina Velikaya",
                "manager_percent": 75.0,
                "ai_percent": 50.0,
                "divergence_count": 1,
                "cells": [
                    {
                        "competence_id": comp_a,
                        "manager_score": 4.0,
                        "ai_score": 2.0,
                        "ai_status": "ready",
                        "divergence": True,
                    },
                    {
                        "competence_id": comp_b,
                        "manager_score": 2.0,
                        "ai_score": None,
                        "ai_status": "not_covered",
                        "divergence": False,
                    },
                ],
            },
            {
                "candidate_vacancy_id": second,
                "name": "Maria Ivanova",
                "manager_percent": None,
                "ai_percent": 100.0,
                "divergence_count": 0,
                "cells": [
                    {
                        "competence_id": comp_a,
                        "manager_score": None,
                        "ai_score": 4.0,
                        "ai_status": "ready",
                        "divergence": False,
                    },
                    {
                        "competence_id": comp_b,
                        "manager_score": None,
                        "ai_score": 4.0,
                        "ai_status": "ready",
                        "divergence": False,
                    },
                ],
            },
        ],
    }


def _csv_rows(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    return [row for row in csv.reader(io.StringIO(text)) if row]


def _xlsx_rows(content: bytes) -> list[list[object]]:
    ws = load_workbook(io.BytesIO(content))["Canvas"]
    # Row 4 is the header; rows 1-3 are the title band.
    return [
        [cell.value for cell in row]
        for row in ws.iter_rows(min_row=4)
        if any(cell.value is not None for cell in row)
    ]


class TestExportMirrorsTheToolbar:
    def test_ai_only_with_a_hidden_candidate_in_both_formats(self) -> None:
        """The ticket's own reproduction: View = AI only, Scale = Percent,
        one candidate unticked. Exactly one AI row must survive, in both
        formats, and both must carry the Total column."""
        matrix = _matrix_stub()
        visible = str(matrix["candidates"][0]["candidate_vacancy_id"])
        built = build_canvas_rows(
            matrix,
            view="ai",
            scale="percent",
            candidate_ids={visible},
        )

        assert built["header"][0] == "Candidate"
        assert built["header"][-1] == "Total"
        assert [row["source"] for row in built["rows"]] == ["AI"]
        assert [row["candidate"] for row in built["rows"]] == ["Ekaterina Velikaya"]

        csv_rows = _csv_rows(render_canvas_csv(matrix, built=built))
        assert csv_rows[0][-1] == "Total"
        assert len(csv_rows) == 2
        body = csv_rows[1]
        assert body[0] == "Ekaterina Velikaya"
        assert body[1] == "AI"
        # 2.0 of 4 with a divergence marker, then the not-covered cell.
        assert body[2] == "50.0% ⚠"
        assert body[3] == "n/a"
        assert body[-1] == "50.0%"

        xlsx_rows = _xlsx_rows(render_canvas_xlsx(matrix, vacancy_title="V", built=built))
        assert xlsx_rows[0][-1] == "Total"
        assert len(xlsx_rows) == 2
        assert xlsx_rows[1][:2] == ["Ekaterina Velikaya", "AI"]
        assert xlsx_rows[1][-1] == "50.0%"
        # The two formats describe the same grid, cell for cell.
        assert [str(v) for v in xlsx_rows[1]] == body

    def test_manager_ai_view_keeps_both_rows_per_candidate(self) -> None:
        built = build_canvas_rows(_matrix_stub())
        assert [row["source"] for row in built["rows"]] == [
            "Manager",
            "AI",
            "Manager",
            "AI",
        ]
        # Only the first row of a candidate repeats the name.
        assert [row["candidate"] for row in built["rows"]] == [
            "Ekaterina Velikaya",
            "",
            "Maria Ivanova",
            "",
        ]

    def test_aggregated_view_collapses_to_one_row(self) -> None:
        built = build_canvas_rows(_matrix_stub(), view="aggregated")
        assert [row["source"] for row in built["rows"]] == [
            "Aggregated",
            "Aggregated",
        ]
        # (4.0 + 2.0) / 2 on the first cell, and the mean of the two
        # percentages in Total.
        assert built["rows"][0]["cells"][0]["text"] == 3.0
        assert built["rows"][0]["total"] == 62.5

    def test_only_divergences_drops_rows_and_columns(self) -> None:
        built = build_canvas_rows(_matrix_stub(), only_divergences=True)
        assert [row["candidate"] for row in built["rows"]] == [
            "Ekaterina Velikaya",
            "",
        ]
        # Communication diverges for nobody left on screen.
        assert built["competences"] == ["Python"]

    def test_hide_unscored_drops_the_candidate_with_no_scores(self) -> None:
        matrix = _matrix_stub()
        for cell in matrix["candidates"][1]["cells"]:
            cell["manager_score"] = None
            cell["ai_score"] = None
            cell["ai_status"] = "missing"
        built = build_canvas_rows(matrix, hide_unscored=True)
        assert {row["candidate"] for row in built["rows"]} == {
            "Ekaterina Velikaya",
            "",
        }

    def test_the_subtitle_names_the_round_slot_not_its_wire_key(self) -> None:
        """"interview_1" is what the filter is called on the wire; the
        reader of the export knows the round as "Interview 1"."""
        matrix = _matrix_stub()
        matrix["round"] = "interview_1"
        matrix["round_slots"] = [
            {"key": "pre_interview", "type": "pre_interview", "number": None},
            {"key": "interview_1", "type": "interview", "number": 1},
        ]
        ws = load_workbook(
            io.BytesIO(render_canvas_xlsx(matrix, vacancy_title="V"))
        )["Canvas"]
        subtitle = str(ws.cell(row=2, column=1).value)
        assert "Round: Interview 1" in subtitle
        assert "interview_1" not in subtitle

    def test_csv_neutralises_a_formula_cell(self) -> None:
        matrix = _matrix_stub()
        matrix["candidates"][0]["name"] = "=cmd|' /C calc'!A0"
        rows = _csv_rows(render_canvas_csv(matrix))
        assert rows[1][0].startswith("'=")


class TestExportEndpoints:
    async def test_both_formats_are_served_and_honour_the_toolbar(
        self, auth_client, db: AsyncSession, tenant, user
    ) -> None:
        vacancy_id, cv_ids = await _vacancy_with_two_candidates(db, tenant, user)
        await db.commit()
        query = (
            f"?view=ai&scale=percent&only_divergences=false&hide_unscored=false"
            f"&candidates={cv_ids[0]}"
        )
        base = f"/api/recruitment/vacancies/{vacancy_id}/assessment-matrix/export"

        csv_res = await auth_client.get(f"{base}.csv{query}")
        assert csv_res.status_code == 200
        assert csv_res.headers["content-type"].startswith("text/csv")
        rows = _csv_rows(csv_res.content)
        assert rows[0][-1] == "Total"
        # One candidate, AI only — one body row.
        assert len(rows) == 2
        assert rows[1][1] == "AI"

        xlsx_res = await auth_client.get(f"{base}.xlsx{query}")
        assert xlsx_res.status_code == 200
        assert len(_xlsx_rows(xlsx_res.content)) == 2


_FRONTEND = Path(__file__).resolve().parents[3] / "frontend"
_CANVAS_PAGE = (
    _FRONTEND
    / "src"
    / "app"
    / "(fullscreen)"
    / "recruitment"
    / "requisitions"
    / "[id]"
    / "assessments"
    / "canvas"
    / "page.tsx"
)


@pytest.mark.skipif(not _CANVAS_PAGE.exists(), reason="frontend tree not present")
class TestExportRouteParity:
    """The export URLs the canvas builds must be URLs the app serves."""

    def test_the_page_calls_urls_that_exist(self) -> None:
        source = _CANVAS_PAGE.read_text(encoding="utf-8")
        template = re.search(
            r"/recruitment/vacancies/\$\{[^}]+\}/assessment-matrix/export\.\$\{[^}]+\}",
            source,
        )
        assert template, "canvas export URL not found — did the page move?"
        mounted = {getattr(route, "path", "") for route in app.routes}
        for suffix in ("xlsx", "csv"):
            assert (
                "/api/recruitment/vacancies/{vacancy_id}"
                f"/assessment-matrix/export.{suffix}" in mounted
            )
        # And the toolbar really travels with them.
        for param in (
            "view",
            "scale",
            "only_divergences",
            "hide_unscored",
            "candidates",
        ):
            assert param in source
